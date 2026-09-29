# -*- coding: utf-8 -*-
"""生成项目报告 Word 文档。"""
import os

import pandas as pd
from docx import Document
from docx.shared import Inches, Pt

BASE = os.path.dirname(os.path.abspath(__file__))
overview = pd.read_csv(os.path.join(BASE, "integrated", "games_overview.csv"))

doc = Document()
doc.add_heading("Steam 游戏生态的多维画像", level=0)
doc.add_paragraph("——多源多维度数据采集、融合分析与可视化")
doc.add_paragraph("《大数据原理与应用》课程中期大作业")

# 一、项目概述
doc.add_heading("一、项目概述与选题", level=1)
doc.add_paragraph(
    "Steam 是全球最大的 PC 游戏分发平台，围绕它聚集了硬件配置、游戏在线人数、玩家评论、"
    "游戏元数据和更新事件等多维数据。本项目围绕「Steam 游戏生态」这一真实问题，从 5 个相互独立"
    "的数据源进行观察与验证，通过实体（appid）、时间、事件三种关联维度将多源数据融合，"
    "在统一的数据基础上完成跨源分析与可视化。"
)

# 二、研究问题
doc.add_heading("二、核心研究问题", level=1)
doc.add_paragraph("本项目提出 3 个核心研究问题，每个均依赖两个及以上数据源：")
problems = [
    ("Q1", "游戏好评率能否解释其在线人气？", "评论 + 在线人数", "实体关联（appid）"),
    ("Q2", "玩家硬件配置是否随新一代游戏发行而协同演进？", "硬件调查 + 游戏元数据", "时间关联"),
    ("Q3", "一次更新/赛事对游戏在线人数能产生多大、多长时间的冲击？", "新闻事件 + 在线人数", "事件关联 + 时间"),
]
t = doc.add_table(rows=1, cols=4)
t.style = "Light Grid Accent 1"
hdr = t.rows[0].cells
for i, h in enumerate(["编号", "研究问题", "依赖数据源", "关联方式"]):
    hdr[i].text = h
for q, p, s, r in problems:
    row = t.add_row().cells
    row[0].text = q
    row[1].text = p
    row[2].text = s
    row[3].text = r

# 三、数据源
doc.add_heading("三、数据源", level=1)
doc.add_paragraph("本项目使用 5 个相互独立的数据源，覆盖「官方统计 + 行为数据 + 事件文本 + 实体映射」4 类形态：")
sources = [
    ("1. Steam 硬件调查", "官方统计（月度）", "31 MB", "2008-11 ~ 2026-08", "date（时间）"),
    ("2. Steam 在线人数", "公开 API（分钟级）", "1.3 MB（扩采 Top100）", "2026-09-17 ~ 19", "appid + timestamp"),
    ("3. Steam 评论", "现成数据集（文本）", "8 游戏 1.83GB", "至 2024", "appid + timestamp"),
    ("4. 游戏元数据", "官方 API（映射）", "辅助源", "2026-09", "appid"),
    ("5. 新闻/更新事件", "官方 API（事件）", "辅助源", "2024-07 ~ 2026-09", "appid + date"),
]
t2 = doc.add_table(rows=1, cols=5)
t2.style = "Light Grid Accent 1"
hdr = t2.rows[0].cells
for i, h in enumerate(["数据源", "类型", "规模", "时间范围", "关联字段"]):
    hdr[i].text = h
for s in sources:
    row = t2.add_row().cells
    for i, v in enumerate(s):
        row[i].text = v

# 四、采集与处理
doc.add_heading("四、数据采集与处理", level=1)
doc.add_paragraph(
    "处理链遵循 Raw Data → Cleaned Data → Integrated Data 三层次。"
    "在线人数数据在采集中发现 123 行 player_count=0 的异常（集中在 09-19 凌晨 00:27~03:06，"
    "为 Steam API 异常返回），已剔除并保留原始备份。"
)
doc.add_paragraph(
    "融合阶段以 appid 为实体键，将元数据、在线人数、评论、新闻四类信息关联为一张"
    "「游戏维度」总表（integrated/games_overview.csv，8 行 × 19 列）。"
)

# 五、跨源关联
doc.add_heading("五、跨源数据关联", level=1)
doc.add_paragraph(
    "本项目使用 3 种跨源关联维度：① 实体关联（appid 统一映射评论文件名、在线 appid、元数据/新闻 appid）；"
    "② 时间关联（硬件月度、在线分钟级、新闻时间戳统一对齐）；③ 事件关联（更新/赛事事件时间点 ↔ 在线人数波动）。"
)
doc.add_picture(os.path.join(BASE, "data_lineage.png"), width=Inches(5.5))

# 六、分析
doc.add_heading("六、数据分析与结果", level=1)
doc.add_heading("Q1 口碑 vs 在线人气（评论 + 在线）", level=2)
top = overview.sort_values("online_mean", ascending=False)
doc.add_paragraph("8 个核心游戏的评论量、好评率与在线人数对比如下：")
t3 = doc.add_table(rows=1, cols=4)
t3.style = "Light Grid Accent 1"
for i, h in enumerate(["游戏", "评论数", "好评率", "在线均值"]):
    t3.rows[0].cells[i].text = h
for _, r in top.iterrows():
    row = t3.add_row().cells
    row[0].text = r["game_name"]
    row[1].text = f"{int(r['review_count']):,}"
    row[2].text = f"{r['positive_rate']*100:.1f}%"
    row[3].text = f"{int(r['online_mean']):,}"
corr = overview["positive_rate"].corr(overview["online_mean"])
doc.add_paragraph(
    f"好评率与在线均值的 Pearson 相关系数 r = {corr:.3f}，几乎无相关。"
    "最典型的是 PUBG：好评率仅 58.2%（8 款中最低），但在线人数第三高（均值 35 万）。"
    "这说明口碑不是在线人气的决定因素，免费制、电竞生态、社交网络等因素作用更大。"
)
doc.add_picture(os.path.join(BASE, "visualization", "v2_reputation_vs_online.png"), width=Inches(5.5))

doc.add_heading("Q2 硬件配置协同演进（硬件 + 元数据）", level=2)
doc.add_paragraph(
    "Steam 硬件调查显示，玩家内存主流从 8GB（2013-2018，占比 26%~47%）切换到 16GB"
    "（2019-2020，占比 40%~50%）。这一代际切换与 Destiny 2（2019）、Apex Legends（2020）"
    "等新一代大型游戏发行时间吻合，说明硬件升级与游戏需求协同演进。"
)
doc.add_picture(os.path.join(BASE, "visualization", "v4_hardware_evolution.png"), width=Inches(5.5))

doc.add_heading("Q3 更新/赛事事件冲击（新闻 + 在线）", level=2)
doc.add_paragraph(
    "对比事件前后 3 小时在线均值：PUBG 电竞赛事（PEC/PAS2）带来 +62%~+68% 的在线增长，"
    "College Royale 活动回归带来 +68.6%；Apex 反作弊更新 +36.4%，CS2 更新 +23.2%。"
    "需注意：凌晨 2 点的事件后在线显示 -44%，实为「日周期」混淆（凌晨在线自然下降），"
    "不能误读为事件负效应——这正是「相关性 ≠ 因果」的体现。"
)
doc.add_picture(os.path.join(BASE, "visualization", "v3_event_impact.png"), width=Inches(5.5))

doc.add_heading("数据挖掘：KMeans 聚类", level=2)
doc.add_paragraph(
    "对 8 个游戏按好评率、在线均值、评论数、平均游玩时长 4 维指标做 KMeans 聚类（3 类）："
    "聚类 0 为「头部巨头」（CS2、Dota2，好评率高、在线极高）；聚类 1 为「口碑好但人气中等」"
    "（Apex、TF2、GTA V、Rust、Destiny 2）；聚类 2 单独为 PUBG——「口碑差但人气高的异类」。"
    "聚类结果印证了 Q1 的结论。"
)

# 七、可视化
doc.add_heading("七、核心可视化", level=1)
doc.add_paragraph("本项目生成 5 张「问题式标题」的核心可视化，覆盖总体特征、跨源关系、时间分析 3 种视角：")
vizs = [
    "v1_online_timeseries.png — 不同游戏的在线人数如何随时间波动？（总体特征）",
    "v2_reputation_vs_online.png — 游戏好评率能解释其在线人气吗？（跨源关系）",
    "v3_event_impact.png — 更新/赛事事件后在线人数是否显著上升？（跨源关系）",
    "v4_hardware_evolution.png — 玩家硬件配置如何随新一代游戏发行而演进？（时间分析）",
    "v5_hourly_rhythm.png — 全球玩家上线高峰出现在哪些时段？（时间分析）",
]
for v in vizs:
    doc.add_paragraph(v, style="List Bullet")

# 八、结论
doc.add_heading("八、主要结论", level=1)
for c in [
    "1. 口碑与在线人气几乎无关（r = -0.068）：免费制、电竞生态、社交网络等因素比口碑更能决定人气。",
    "2. 硬件配置与游戏发行协同演进：内存主流从 8GB 切换到 16GB（2019-2020），与新一批大型游戏发行吻合。",
    "3. 更新/赛事显著拉动在线：电竞与活动带来 +60% 以上的在线增长，但需剔除日周期混淆。",
]:
    doc.add_paragraph(c)

# 九、局限
doc.add_heading("九、局限性与数据偏差", level=1)
for c in [
    "1. 核心研究对象仅 8 个游戏，样本量小，统计结论（如相关性）仅供探索性参考。",
    "2. 在线人数为行为数据，规模受采集周期限制（分钟级需持续积累）。",
    "3. 评论数据规模占主导（1.83GB），其他源为官方统计/映射/事件等天然小规模源。",
    "4. 新闻数据仅 100 条/游戏，无法覆盖更长历史的事件。",
]:
    doc.add_paragraph(c)

# 十、AI 使用
doc.add_heading("十、AI 使用说明", level=1)
doc.add_paragraph(
    "本项目允许并合理使用 AI 编程助手（Cline/Claude），参与了数据采集、清洗、融合、分析和可视化等环节。"
    "所有 AI 生成的代码均在本地实际运行验证，数据结论经人工核验（详见 AI_USAGE.md）。"
)

out = os.path.join(BASE, "项目报告_Steam游戏生态.docx")
doc.save(out)
print(f"报告已生成: {out}")
